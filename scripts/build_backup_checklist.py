from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Docs" / "Complete_Backup_and_Recovery_Checklist.md"
OUTPUT = ROOT / "Docs" / "Complete_Backup_and_Recovery_Checklist.docx"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 99, 110)


def set_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for list_style_name in ("List Bullet", "List Number"):
        list_style = doc.styles[list_style_name]
        list_style.font.name = "Calibri"
        list_style.font.size = Pt(11)
        list_style.paragraph_format.left_indent = Inches(0.375)
        list_style.paragraph_format.first_line_indent = Inches(-0.188)
        list_style.paragraph_format.space_after = Pt(4)
        list_style.paragraph_format.line_spacing = 1.25
        list_style.paragraph_format.keep_together = True

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = "TUDELA GLOBAL  |  BACKUP AND RECOVERY"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.runs[0], size=8.5, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("3Sci  |  Living recovery checklist"), size=8.5, color=MUTED)


def add_inline(paragraph, text):
    cursor = 0
    while cursor < len(text):
        bold_at = text.find("**", cursor)
        code_at = text.find("`", cursor)
        candidates = [(pos, kind) for pos, kind in ((bold_at, "bold"), (code_at, "code")) if pos >= 0]
        if not candidates:
            set_font(paragraph.add_run(text[cursor:]))
            break
        pos, kind = min(candidates)
        if pos > cursor:
            set_font(paragraph.add_run(text[cursor:pos]))
        marker = "**" if kind == "bold" else "`"
        end = text.find(marker, pos + len(marker))
        if end < 0:
            set_font(paragraph.add_run(text[pos:]))
            break
        run = paragraph.add_run(text[pos + len(marker):end])
        set_font(run, bold=(kind == "bold"), color=DARK_BLUE if kind == "code" else None)
        if kind == "code":
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Consolas")
        cursor = end + len(marker)


def add_manual_list_paragraph(doc, marker, text):
    """Create a stable hanging-indent list item that LibreOffice renders reliably."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.keep_together = True
    set_font(paragraph.add_run(f"{marker}\t"), size=11)
    add_inline(paragraph, text)
    return paragraph


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    set_font(title.add_run("Tudela Global"), size=28, bold=True, color=DARK_BLUE)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    set_font(subtitle.add_run("Complete Backup and Recovery Checklist"), size=17, bold=True, color=BLUE)

    index = 1
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            index += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
            index += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
            index += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(9)
            add_inline(p, line[2:])
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = DARK_BLUE
            index += 1
            continue
        if line.startswith("| "):
            block = []
            while index < len(lines) and lines[index].startswith("|"):
                block.append(lines[index])
                index += 1
            rows = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in block]
            if len(rows) > 1 and all(set(cell) <= {"-", ":"} for cell in rows[1]):
                rows.pop(1)
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            width_patterns = {2: [2700, 6660], 3: [2300, 4060, 3000], 4: [1800, 3000, 2760, 1800]}
            set_table_geometry(table, width_patterns.get(len(rows[0]), [9360 // len(rows[0])] * len(rows[0])))
            for r, values in enumerate(rows):
                row_properties = table.rows[r]._tr.get_or_add_trPr()
                cant_split = OxmlElement("w:cantSplit")
                row_properties.append(cant_split)
                if r == 0:
                    repeat_header = OxmlElement("w:tblHeader")
                    repeat_header.set(qn("w:val"), "true")
                    row_properties.append(repeat_header)
                for c, value in enumerate(values):
                    cell = table.cell(r, c)
                    cell.text = ""
                    add_inline(cell.paragraphs[0], value)
                    for run in cell.paragraphs[0].runs:
                        set_font(run, size=9.5, bold=(r == 0), color=DARK_BLUE if r == 0 else None)
                    if r == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "E8EEF5")
                        cell._tc.get_or_add_tcPr().append(shading)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if line.startswith("- ["):
            add_manual_list_paragraph(doc, "☒" if line.startswith("- [x] ") else "☐", line[6:])
            index += 1
            continue
        if line.startswith("- "):
            add_manual_list_paragraph(doc, "•", line[2:])
            index += 1
            continue
        if re.match(r"^\d+\.\s", line):
            number, text = line.split(". ", 1)
            add_manual_list_paragraph(doc, f"{number}.", text)
            index += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, line.replace("  ", " "))
        index += 1

    doc.core_properties.title = "Tudela Global Complete Backup and Recovery Checklist"
    doc.core_properties.subject = "Programs, cloud services, backup set and recovery procedure"
    doc.core_properties.author = "3Sci / Tudela Global"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
