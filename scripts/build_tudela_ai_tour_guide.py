from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PROMPT_PATH = ROOT / "Docs" / "Tudela_AI_Tour_Creation_Master_Prompt.txt"
OUTPUT_PATH = ROOT / "Docs" / "Tudela_AI_Tour_Creation_Master_Instructions.docx"

NAVY = "0B2559"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GOLD = "D6A400"
INK = "172033"
MUTED = "596579"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F7FB"
LIGHT_GOLD = "FFF7D6"
WHITE = "FFFFFF"
BORDER = "C8D2E2"


def set_run_font(run, size=None, color=INK, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def shade_paragraph(paragraph, fill, border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        for edge in ("top", "left", "bottom", "right"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), "6")
            node.set(qn("w:space"), "4")
            node.set(qn("w:color"), border)
            p_bdr.append(node)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "left" if edge == "start" else "right" if edge == "end" else edge
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])
    set_run_font(run, size=9, color=MUTED)


def configure_numbering(doc):
    numbering = doc.part.numbering_part.element
    max_abs = max([int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))] or [0])
    max_num = max([int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))] or [0])

    def add_definition(kind, fmt, text, bullet_font=None):
        nonlocal max_abs, max_num
        max_abs += 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(max_abs))
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "space" if kind == "bullet" else "tab")
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        level.extend([start, num_fmt, lvl_text, suffix, lvl_jc, p_pr])
        if bullet_font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), bullet_font)
            r_fonts.set(qn("w:hAnsi"), bullet_font)
            r_pr.append(r_fonts)
            level.append(r_pr)
        abstract.append(level)
        numbering.append(abstract)
        max_num += 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(max_num))
        abs_id = OxmlElement("w:abstractNumId")
        abs_id.set(qn("w:val"), str(max_abs))
        num.append(abs_id)
        numbering.append(num)
        return max_num

    return add_definition("bullet", "bullet", "•", "Calibri"), add_definition("decimal", "decimal", "%1.")


def configure_builtin_bullet(doc):
    style = doc.styles["List Bullet"]
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(11)
    style.paragraph_format.left_indent = Inches(0.375)
    style.paragraph_format.first_line_indent = Inches(-0.188)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25

    num_id = style._element.pPr.numPr.numId.val
    numbering = doc.part.numbering_part.element
    num = next(node for node in numbering.findall(qn("w:num")) if int(node.get(qn("w:numId"))) == num_id)
    abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = next(node for node in numbering.findall(qn("w:abstractNum")) if node.get(qn("w:abstractNumId")) == abstract_id)
    level = abstract.find(qn("w:lvl"))
    level.find(qn("w:lvlText")).set(qn("w:val"), "•")
    r_pr = level.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        level.append(r_pr)
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    suffix = level.find(qn("w:suff"))
    if suffix is None:
        suffix = OxmlElement("w:suff")
        level.insert(3, suffix)
    suffix.set(qn("w:val"), "space")
    p_pr = level.find(qn("w:pPr"))
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is not None:
        p_pr.remove(tabs)
    ind = p_pr.find(qn("w:ind"))
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)


def new_num_instance(doc, template_num_id):
    numbering = doc.part.numbering_part.element
    template = next(node for node in numbering.findall(qn("w:num")) if node.get(qn("w:numId")) == str(template_num_id))
    abstract_id = template.find(qn("w:abstractNumId")).get(qn("w:val"))
    max_num = max([int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))] or [0])
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(max_num + 1))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return max_num + 1


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_heading(doc, text, level):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(paragraph)
    return paragraph


def add_callout(doc, title, body, fill=LIGHT_GOLD):
    p = doc.add_paragraph()
    shade_paragraph(p, fill, BORDER)
    p.paragraph_format.left_indent = Pt(7)
    p.paragraph_format.right_indent = Pt(7)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(title + "\n")
    set_run_font(r, size=11, color=NAVY, bold=True)
    r2 = p.add_run(body)
    set_run_font(r2, size=10.5, color=INK)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25

    tokens = {
        1: (16, BLUE, 18, 10),
        2: (13, BLUE, 14, 7),
        3: (12, DARK_BLUE, 10, 5),
    }
    for level, (size, color, before, after) in tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    configure_builtin_bullet(doc)


def add_intro(doc, bullet_num, decimal_num):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(0)
    r = kicker.add_run("TUDELA GLOBAL CONTENT STANDARD")
    set_run_font(r, size=10, color=GOLD, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(7)
    r = title.add_run("Universal AI Tour-Creation Instructions")
    set_run_font(r, size=28, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run("A copy-ready master prompt for consistent, researched, five-stop Tudela Global tours")
    set_run_font(r, size=13.5, color=MUTED)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 6660])
    for index, value in enumerate(("Document detail", "Value")):
        shade_cell(table.rows[0].cells[index], LIGHT_BLUE)
        p = table.rows[0].cells[index].paragraphs[0]
        r = p.add_run(value)
        set_run_font(r, size=9.5, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    values = [
        ("Prepared for", "3Sci / Tudela Global"),
        ("Standard", "Version 1.0 — 10 August 2026"),
        ("Primary output", "One CMS-ready tour with exactly five ordered stops"),
        ("Publication state", "Draft only — CMS review and approval required"),
    ]
    for label, value in values:
        cells = table.add_row().cells
        for index, text in enumerate((label, value)):
            set_cell_width(cells[index], (2700, 6660)[index])
            set_cell_margins(cells[index])
            p = cells[index].paragraphs[0]
            r = p.add_run(text)
            set_run_font(r, size=10, color=INK, bold=index == 0)

    doc.add_paragraph()
    add_callout(
        doc,
        "Important",
        "This guide standardizes the content an AI produces. It does not bypass Tudela’s editorial workflow. The generated package must be checked, saved as a draft, reviewed, approved, and published by an authorized CMS user.",
    )

    add_heading(doc, "1. Purpose", 1)
    doc.add_paragraph(
        "Use this standard when asking any capable AI system to research and write a complete Tudela Global tour. It converts the source document-assisted import parameters into a strict content contract: five real stops, exact CMS field labels, verified locations, traceable sources, lawful image candidates, and a final validation report."
    )
    doc.add_paragraph(
        "The result is suitable for structured CMS entry and for the planned document-assisted importer. At the current development stage, an authorized editor must still review and enter or import the content through the available CMS workflow."
    )

    add_heading(doc, "2. How to use this guide", 1)
    steps = [
        "Complete the request card below. City, country, and tour theme are mandatory.",
        "Copy the entire master prompt from BEGIN PROMPT through END PROMPT. For easiest copying, use the companion plain-text file supplied with this guide.",
        "Paste it into the chosen AI channel, followed by the completed request card.",
        "If the AI cannot browse, map, or verify image rights, accept only a package clearly marked REQUIRES HUMAN VERIFICATION. Do not ask it to guess.",
        "Review the returned Validation Report before transferring any content into Tudela. Publication still requires the normal editor-to-reviewer approval process.",
    ]
    for item in steps:
        p = doc.add_paragraph(item)
        apply_num(p, decimal_num)

    add_heading(doc, "3. Request card to send with the prompt", 1)
    request_lines = [
        "City / Destination Name: [required]",
        "Country: [required]",
        "Tour Theme or Story Focus: [required]",
        "Target Audience: [optional]",
        "Preferred Starting Point: [optional]",
        "Maximum Distance or Duration: [optional]",
        "Accessibility Requirements: [optional]",
        "Access Model: [optional: Free, Paid, or Subscription]",
        "Source Language: [optional: English, Hebrew, Russian, French, German, or Spanish]",
        "Requested Translations: [optional]",
        "Required or Excluded Stops: [optional]",
        "Special Editorial or Cultural Requirements: [optional]",
    ]
    table = doc.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    header_cell = table.cell(0, 0)
    shade_cell(header_cell, LIGHT_BLUE)
    p = header_cell.paragraphs[0]
    r = p.add_run("Copy and complete this request card")
    set_run_font(r, size=10, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    cell = table.cell(1, 0)
    shade_cell(cell, PALE_BLUE)
    first = True
    for line in request_lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(2)
        label, rest = line.split(":", 1)
        r1 = p.add_run(label + ":")
        set_run_font(r1, size=10, color=NAVY, bold=True)
        r2 = p.add_run(rest)
        set_run_font(r2, size=10, color=MUTED)

    add_heading(doc, "4. What a conforming AI response contains", 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    widths = [2700, 6660]
    set_table_geometry(table, widths)
    headers = ["Component", "Required result"]
    for idx, value in enumerate(headers):
        shade_cell(table.rows[0].cells[idx], LIGHT_BLUE)
        p = table.rows[0].cells[idx].paragraphs[0]
        r = p.add_run(value)
        set_run_font(r, size=10, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("CMS Import Package", "City, tour, and exactly five numbered stops using exact field labels."),
        ("Route", "Verified real-world coordinates, practical order, distance, duration, and walking directions."),
        ("Visitor content", "Overview, descriptions, full narratives, and natural audio transcripts."),
        ("Research", "A traceable Sources Register linked to the tour and each stop."),
        ("Media", "One rights-checked image candidate per stop plus an optional hero candidate."),
        ("Quality control", "A field-by-field Validation Report with unresolved work disclosed."),
    ]
    for left, right in rows:
        cells = table.add_row().cells
        for index, value in enumerate((left, right)):
            set_cell_width(cells[index], widths[index])
            set_cell_margins(cells[index])
            p = cells[index].paragraphs[0]
            r = p.add_run(value)
            set_run_font(r, size=9.5, color=INK, bold=index == 0)

    add_heading(doc, "5. Acceptance gates", 1)
    gates = [
        "Exactly five real, non-duplicate stops in a logical walking order.",
        "No fabricated facts, coordinates, addresses, quotations, access claims, or image rights.",
        "Only approved Tudela Interest values.",
        "All required CMS fields present, with uncertainty explicitly flagged.",
        "Traceable research sources and image rights records.",
        "No publication until an authorized reviewer has approved the draft in the CMS.",
    ]
    for item in gates:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_page_break()


def add_prompt(doc, prompt_text, bullet_num, decimal_num):
    add_heading(doc, "Copy-ready master prompt", 1)
    p = doc.add_paragraph()
    r = p.add_run("Copy from BEGIN PROMPT through END PROMPT without deleting or renaming fields.")
    set_run_font(r, size=10.5, color=NAVY, bold=True)

    section_headings = {
        "BEGIN PROMPT",
        "YOUR TASK",
        "REQUIRED INPUTS",
        "DEFAULTS WHEN OPTIONAL INPUTS ARE NOT SUPPLIED",
        "NON-NEGOTIABLE ACCURACY RULES",
        "GPS AND ROUTE RULES",
        "STOP-SELECTION RULES",
        "WRITING AND LENGTH STANDARD",
        "NARRATIVE STANDARD FOR EVERY STOP",
        "IMAGE AND RIGHTS STANDARD",
        "LANGUAGE RULES",
        "OUTPUT RULES",
        "EXACT OUTPUT STRUCTURE",
        "TUDELA GLOBAL TOUR PACKAGE",
        "CMS IMPORT PACKAGE",
        "TRANSLATIONS",
        "SOURCES REGISTER",
        "MEDIA RIGHTS MANIFEST",
        "VALIDATION REPORT",
        "FINAL SELF-CHECK BEFORE RESPONDING",
        "END PROMPT",
    }
    stop_heading_re = re.compile(r"^STOP [1-5](?: TRANSLATION)?$")
    translation_heading_re = re.compile(r"^TRANSLATION — ")
    numbered_re = re.compile(r"^(\d+)\.\s+(.*)$")
    prompt_decimal_num = decimal_num

    lines = prompt_text.splitlines()
    for line in lines:
        text = line.rstrip()
        if not text:
            continue
        if text.startswith("TUDELA GLOBAL — MASTER AI TOUR-CREATION PROMPT") or text.startswith("Version 1.0 |"):
            continue
        if text in {"BEGIN PROMPT", "END PROMPT"}:
            p = doc.add_paragraph()
            shade_paragraph(p, NAVY)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(text)
            set_run_font(r, size=11, color=WHITE, bold=True)
            continue
        if text in section_headings:
            add_heading(doc, text.title() if text != "CMS IMPORT PACKAGE" else "CMS Import Package", 2)
            continue
        if stop_heading_re.match(text) or translation_heading_re.match(text):
            add_heading(doc, text, 3)
            continue
        if text == "TOUR":
            add_heading(doc, "TOUR", 3)
            continue
        match = numbered_re.match(text)
        if match:
            if match.group(1) == "1":
                prompt_decimal_num = new_num_instance(doc, decimal_num)
            p = doc.add_paragraph(match.group(2))
            apply_num(p, prompt_decimal_num)
            continue
        if text.lstrip().startswith("- "):
            doc.add_paragraph(text.lstrip()[2:], style="List Bullet")
            continue
        p = doc.add_paragraph()
        p.paragraph_format.widow_control = True
        if ":" in text and not text.startswith("http"):
            label, value = text.split(":", 1)
            if len(label) <= 65 and not label.startswith("["):
                r1 = p.add_run(label + ":")
                set_run_font(r1, size=10.5, color=NAVY, bold=True)
                r2 = p.add_run(value)
                set_run_font(r2, size=10.5, color=INK)
                continue
        r = p.add_run(text)
        if text.startswith("[") and text.endswith("]"):
            set_run_font(r, size=10, color=MUTED, italic=True)
        else:
            set_run_font(r, size=10.5, color=INK)


def build():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r1 = hp.add_run("TUDELA GLOBAL")
    set_run_font(r1, size=8.5, color=NAVY, bold=True)
    r2 = hp.add_run("  |  Master AI Tour-Creation Standard")
    set_run_font(r2, size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run("3Sci / Tudela Global  •  Version 1.0  •  Page ")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_number(fp)

    bullet_num, decimal_num = configure_numbering(doc)
    add_intro(doc, bullet_num, decimal_num)
    add_prompt(doc, PROMPT_PATH.read_text(encoding="utf-8"), bullet_num, decimal_num)

    props = doc.core_properties
    props.title = "Tudela Global Universal AI Tour-Creation Instructions"
    props.subject = "Unified five-stop tour generation and CMS content standard"
    props.author = "3Sci / Tudela Global"
    props.keywords = "Tudela Global, CMS, AI prompt, heritage tour, five stops, GPS, image rights"
    props.comments = "Generated from the Tudela document-assisted import parameters and current CMS schema."

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build()
