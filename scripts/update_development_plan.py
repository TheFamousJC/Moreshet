from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
PLAN = ROOT / "Docs" / "Codex Production Development Plan.docx"


def insert_before(anchor: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def main() -> None:
    document = Document(PLAN)
    import_checkpoint = "Stage 5 checkpoint - Manual MP3 and WAV narration import"
    acceptance_checkpoint = "Stage 5 acceptance - Published Kokoro narration"
    existing = {p.text.strip() for p in document.paragraphs}
    if import_checkpoint in existing and acceptance_checkpoint in existing:
        print(f"Development plan already current: {PLAN}")
        return

    offline_heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() == "7. Offline operation"
    )
    additions = []
    if import_checkpoint not in existing:
        additions.extend((
            (import_checkpoint, "Heading 3"),
            ("Implemented 13 August 2026 as a cost-control and provider-independence path.", "Normal"),
            ("Editors can upload English or multilingual MP3 and WAV narration produced by Kokoro, a human narrator or another licensed TTS system.", "Normal"),
            ("The server validates organization access, stop ownership, saved transcript version, language, file size, audio structure and checksum before private storage.", "Normal"),
            ("Every import records its source, model, narrator or voice and original filename, and follows the existing submit, review, approve and publish workflow.", "Normal"),
            ("Manual imports bypass paid TTS generation and create no generation-usage event; review, storage, delivery, licensing and quality controls still apply.", "Normal"),
        ))
    if acceptance_checkpoint not in existing:
        additions.extend((
            (acceptance_checkpoint, "Heading 3"),
            ("Completed 13 August 2026 with a real 51-second English Kokoro MP3 for Umm el-Umdan Ancient Synagogue.", "Normal"),
            ("Private preview, review submission, approval, publication, signed public delivery and owner-confirmed visitor playback all passed.", "Normal"),
            ("The stored recording can be replayed without regenerating speech and did not invoke paid Google Text-to-Speech.", "Normal"),
            ("Detailed pronunciation, WAV-specific, Hebrew/native-speaker, other-language and failure-mode acceptance remain pending.", "Normal"),
        ))
    for text, style in additions:
        insert_before(offline_heading, text, style)

    document.save(PLAN)
    print(PLAN)


if __name__ == "__main__":
    main()
