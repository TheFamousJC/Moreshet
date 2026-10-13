from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Docs" / "Tudela_Global_User_Manual_and_Technical_Reference.md"
OUTPUT = ROOT / "Docs" / "Tudela_Global_User_Manual_and_Technical_Reference.docx"


def source_metadata(label: str, default: str) -> str:
    source_text = SOURCE.read_text(encoding="utf-8")
    match = re.search(rf"^\*\*{re.escape(label)}:\*\*\s*(.+?)\s*$", source_text, re.MULTILINE)
    return match.group(1).strip() if match else default


DOCUMENT_VERSION = source_metadata("Document version", "Unversioned")
PLATFORM_STAGE = source_metadata("Platform stage", "Living development")
LAST_UPDATED = source_metadata("Last updated", "Not recorded")
STAGING_ADDRESS = source_metadata("Staging address", "Not configured")

NAVY = "17365D"
BLUE = "2E74B5"
MID_BLUE = "1F4D78"
GOLD = "C59A36"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "F7F1E2"
LIGHT_GREY = "F3F5F7"
DARK = "263238"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_width(table, width_dxa=9360) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)


def set_repeat_header_text(section, text: str) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.text = text
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.style = "Header"
    run = p.runs[0]
    run.font.name = "Calibri"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MID_BLUE)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), GOLD)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char, instr_text, fld_char_end))


def configure_section(section, first=False) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.86 if first else 0.82)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.38)
    section.different_first_page_header_footer = first
    set_repeat_header_text(section, "TUDELA GLOBAL  |  USER MANUAL & TECHNICAL REFERENCE")
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = "Footer"
    r = p.add_run(f"Living development manual  •  Version {DOCUMENT_VERSION}  •  Page ")
    r.font.name = "Calibri"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MID_BLUE)
    add_page_number(p)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 10),
        ("Subtitle", 15, MID_BLUE, 0, 18),
        ("Heading 1", 17, BLUE, 18, 9),
        ("Heading 2", 13.5, MID_BLUE, 13, 6),
        ("Heading 3", 11.5, NAVY, 10, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.15

    code = styles.add_style("Manual Code", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string(DARK)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.08)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(5)
    code.paragraph_format.keep_together = True

    note = styles.add_style("Manual Note", WD_STYLE_TYPE.PARAGRAPH)
    note.font.name = "Calibri"
    note.font.size = Pt(9.5)
    note.font.italic = True
    note.font.color.rgb = RGBColor.from_string(NAVY)
    note.paragraph_format.left_indent = Inches(0.22)
    note.paragraph_format.right_indent = Inches(0.12)
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(7)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\)|https?://\S+)")


def add_inline(paragraph, text: str, *, bold=False, color=None) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            r = paragraph.add_run(text[pos:match.start()])
            r.bold = bold
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        token = match.group(0)
        if token.startswith("**"):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        elif token.startswith("`"):
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor.from_string(MID_BLUE)
        elif token.startswith("["):
            label = token[1:token.index("]")]
            url = token[token.index("(") + 1:-1]
            r = paragraph.add_run(f"{label} ({url})")
            r.font.color.rgb = RGBColor.from_string(BLUE)
            r.font.underline = True
        else:
            r = paragraph.add_run(token)
            r.font.color.rgb = RGBColor.from_string(BLUE)
        pos = match.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.bold = bold
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def shade_paragraph(paragraph, fill: str, border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if border:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "16")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border)
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_rule(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), GOLD)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def is_table_separator(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def split_table_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    usable = Inches(6.5)
    width = usable / columns
    for row_idx, data in enumerate(rows):
        row = table.rows[row_idx]
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        if row_idx == 0:
            set_repeat_table_header(row)
        for col_idx, cell in enumerate(row.cells):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if col_idx < len(data):
                add_inline(p, data[col_idx], bold=row_idx == 0)
            for run in p.runs:
                run.font.size = Pt(8.5 if columns >= 3 else 9)
                if row_idx == 0:
                    run.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("TUDELA")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(GOLD)
    run = p.add_run("  GLOBAL")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(28)
    p.add_run("User Manual\nand Technical Reference")

    p = doc.add_paragraph(style="Subtitle")
    p.add_run("A living guide to the visitor application, CMS, GPS, AI and platform foundation")
    add_rule(doc)

    details = [
        ("DOCUMENT VERSION", DOCUMENT_VERSION),
        ("PLATFORM STAGE", PLATFORM_STAGE),
        ("LAST UPDATED", LAST_UPDATED),
        ("STAGING ADDRESS", STAGING_ADDRESS),
    ]
    table = doc.add_table(rows=len(details), cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_repeat_table_header(table.rows[0])
    for row, (label, value) in zip(table.rows, details):
        row.cells[0].width = Inches(1.85)
        row.cells[1].width = Inches(4.4)
        for cell in row.cells:
            set_cell_margins(cell, 70, 80, 70, 80)
        set_cell_shading(row.cells[0], PALE_BLUE)
        p0 = row.cells[0].paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(8.5)
        r0.font.color.rgb = RGBColor.from_string(MID_BLUE)
        p1 = row.cells[1].paragraphs[0]
        r1 = p1.add_run(value)
        r1.font.size = Pt(9.5)
    p = doc.add_paragraph(style="Manual Note")
    p.paragraph_format.space_before = Pt(24)
    add_inline(p, "This edition records the functionality actually implemented and verified. Planned production features are clearly separated from working features.")
    shade_paragraph(p, PALE_GOLD, GOLD)
    p.add_run().add_break(WD_BREAK.PAGE)

    doc.add_heading("Contents", level=1)
    for item in (
        "Document Control",
        "Part I — User Manual",
        "1. Current Platform Status",
        "2. Starting and Stopping the Platform",
        "3–6. Visitor, GPS, CMS and Troubleshooting",
        "Part II — Technical Reference",
        "7–21. Architecture, Services, Data, Security and QA",
        "Development Stage Register",
        "Future Stage Section Template and Change Log",
    ):
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, item)
    p = doc.add_paragraph(style="Manual Note")
    add_inline(p, "Use Word's Navigation pane to move through the complete heading structure.")
    shade_paragraph(p, LIGHT_GREY, BLUE)
    p.add_run().add_break(WD_BREAK.PAGE)


def parse_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    # The source cover ends at its first horizontal rule; Word has a designed cover.
    try:
        start = lines.index("---") + 1
    except ValueError:
        start = 0
    lines = lines[start:]
    i = 0
    in_code = False
    code_lines: list[str] = []
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            p = doc.add_paragraph()
            add_inline(p, " ".join(s.strip() for s in paragraph_buffer))
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                p = doc.add_paragraph(style="Manual Code")
                p.add_run("\n".join(code_lines))
                shade_paragraph(p, LIGHT_GREY, BLUE)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not stripped:
            flush_paragraph()
            i += 1
            continue
        if stripped == "---":
            flush_paragraph()
            add_rule(doc)
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            flush_paragraph()
            rows = [split_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if level == 1 and (title.startswith("Part ") or title in {"Development Stage Register", "Local Reference Documents"}):
                if len(doc.paragraphs) > 2:
                    doc.add_page_break()
            p = doc.add_heading(level=level)
            add_inline(p, title)
            i += 1
            continue
        if stripped.startswith(">"):
            flush_paragraph()
            p = doc.add_paragraph(style="Manual Note")
            add_inline(p, stripped[1:].strip())
            shade_paragraph(p, PALE_GOLD, GOLD)
            i += 1
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
            i += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            flush_paragraph()
            p = doc.add_paragraph(style="List Number")
            add_inline(p, numbered.group(1))
            i += 1
            continue
        paragraph_buffer.append(stripped.rstrip(" "))
        i += 1
    flush_paragraph()


def add_document_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "Tudela Global — User Manual and Technical Reference"
    props.subject = "Maintained product and technical manual through Stage 7 development closure"
    props.author = "Tudela Global"
    props.keywords = "Tudela Global, user manual, CMS, GPS, AI, Supabase, technical reference"
    props.comments = "Regenerate after every completed development stage."


def main(output: Path = OUTPUT) -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], first=True)
    add_document_properties(doc)
    add_cover(doc)
    parse_markdown(doc, markdown)
    # Ensure table rows do not split across pages where possible.
    for table in doc.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build the Tudela Global living manual.")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    main(args.output)
